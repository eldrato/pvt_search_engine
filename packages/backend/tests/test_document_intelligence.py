import os
import io
import pytest
import pymupdf
from docx import Document as DocxDoc
from httpx import AsyncClient, ASGITransport

from app.main import app
from app.ingestion.parsers import PDFParser, DocxParser, CodeParser, TextParser, get_parser
from app.ingestion.chunker import Chunker
from app.services.ingestion import ingestion_service
from app.services.document_intelligence import document_intelligence_service
from app.core.db import get_async_session


@pytest.fixture
def sample_pdf_bytes() -> bytes:
    doc = pymupdf.open()
    # Page 1
    page1 = doc.new_page(width=595, height=842)
    page1.insert_text((50, 72), "Aegis Document Intelligence Specification", fontsize=18)
    page1.insert_text((50, 110), "Section 1: Exact location tracking enables precision navigation.", fontsize=11)
    page1.insert_text((50, 140), "Every assertion requires verifiable source spans and page anchors.", fontsize=11)
    # Page 2
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 72), "Section 2: Highlighting & Visual Evidence", fontsize=16)
    page2.insert_text((50, 110), "Inline PDF highlighter renders directly on the canvas without latency.", fontsize=11)
    
    pdf_bytes = doc.tobytes()
    doc.close()
    return pdf_bytes


@pytest.fixture
def sample_docx_bytes() -> bytes:
    doc = DocxDoc()
    doc.add_heading("System Architecture Document", level=1)
    doc.add_paragraph("First paragraph describing the single source of truth database model.")
    doc.add_heading("Retrieval Engine", level=2)
    doc.add_paragraph("Second paragraph detailing reciprocal rank fusion across BM25 and vector scores.")
    
    bio = io.BytesIO()
    doc.save(bio)
    return bio.getvalue()


def test_pdf_parser_bounding_boxes_and_paragraphs(sample_pdf_bytes):
    parser = PDFParser()
    parsed = parser.parse(sample_pdf_bytes, "intelligence_spec.pdf")

    assert parsed.title == "Intelligence Spec" or "Specification" in parsed.title
    assert parsed.mime_type == "application/pdf"
    assert parsed.metadata["total_pages"] == 2
    assert len(parsed.blocks) >= 2

    # Verify first block has page, paragraph index, and bounding box coordinates
    b0 = parsed.blocks[0]
    assert b0.page_number == 1
    assert b0.paragraph_index is not None
    assert b0.bbox is not None
    assert len(b0.bbox) == 4
    # x0, y0, x1, y1
    assert b0.bbox[0] < b0.bbox[2]
    assert b0.bbox[1] < b0.bbox[3]
    assert "page_width" in b0.extra_meta
    assert "page_height" in b0.extra_meta


def test_chunker_precision_anchors_and_location_meta(sample_pdf_bytes):
    parser = PDFParser()
    parsed = parser.parse(sample_pdf_bytes, "intelligence_spec.pdf")

    # 1. Single page chunk with target_chars small enough to keep on page 1
    small_chunker = Chunker(target_chars=60, overlap_chars=10)
    single_page_chunks = small_chunker.chunk_document(parsed)
    assert len(single_page_chunks) >= 3
    c_single = single_page_chunks[0]
    assert c_single.location_meta["pages"] == [1]
    assert "Page 1" in c_single.anchor_label
    assert "¶" in c_single.anchor_label
    assert len(c_single.location_meta["bboxes"]) > 0

    # 2. Multi-page chunk with target_chars spanning page 1 and page 2
    multi_chunker = Chunker(target_chars=220, overlap_chars=20)
    multi_chunks = multi_chunker.chunk_document(parsed)
    assert len(multi_chunks) >= 1
    c_multi = multi_chunks[0]
    assert len(c_multi.location_meta["pages"]) == 2
    assert "Pages 1-2" in c_multi.anchor_label


def test_docx_parser_paragraph_tracking(sample_docx_bytes):
    parser = DocxParser()
    parsed = parser.parse(sample_docx_bytes, "system_arch.docx")

    assert parsed.mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert len(parsed.blocks) >= 2
    for b in parsed.blocks:
        assert b.paragraph_index is not None
        assert b.section_title is not None
        assert "section" in b.extra_meta


def test_code_parser_line_anchors():
    code = b"""import os
from typing import Dict

def render_preview(path: str) -> Dict[str, str]:
    if not os.path.exists(path):
        return {"error": "not found"}
    return {"status": "ok"}
"""
    parser = CodeParser()
    parsed = parser.parse(code, "viewer.py")

    assert parsed.metadata["language"] == "python"
    assert len(parsed.blocks) >= 1
    b0 = parsed.blocks[0]
    assert b0.line_start == 1
    assert b0.line_end >= 7
    assert "viewer.py:1-" in b0.section_title


@pytest.mark.asyncio
async def test_document_intelligence_pdf_rendering_and_context(sample_pdf_bytes):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as ac:
        # 1. Ingest PDF
        files = {
            "file": ("phase5_test_doc.pdf", sample_pdf_bytes, "application/pdf")
        }
        res = await ac.post("/api/v1/documents/upload", files=files, data={"source_type": "upload"})
        assert res.status_code == 201
        data = res.json()
        doc_id = data["document"]["id"]

        # 2. Test Get Original File
        file_res = await ac.get(f"/api/v1/documents/{doc_id}/file")
        assert file_res.status_code == 200
        assert file_res.headers["content-type"].startswith("application/pdf")
        assert file_res.content.startswith(b"%PDF")

        # 3. Test Render PDF Page as PNG with Highlights
        page_res = await ac.get(
            f"/api/v1/documents/{doc_id}/pdf/page/1?highlight_query=Intelligence&highlight_chunk=0"
        )
        assert page_res.status_code == 200
        assert page_res.headers["content-type"] == "image/png"
        assert page_res.content.startswith(b"\x89PNG\r\n\x1a\n")

        # 4. Test Annotated PDF Generation
        annot_res = await ac.get(
            f"/api/v1/documents/{doc_id}/pdf/highlighted?highlight_query=Intelligence"
        )
        assert annot_res.status_code == 200
        assert annot_res.headers["content-type"] == "application/pdf"
        assert annot_res.content.startswith(b"%PDF")

        # 5. Test Document Structural Locations Map
        loc_res = await ac.get(f"/api/v1/documents/{doc_id}/locations")
        assert loc_res.status_code == 200
        loc_data = loc_res.json()
        assert loc_data["document_id"] == doc_id
        assert loc_data["is_pdf"] is True
        assert loc_data["total_pages"] == 2
        assert len(loc_data["chunks"]) > 0
        assert "anchor_label" in loc_data["chunks"][0]

        # 6. Test Chunk Location Context
        ctx_res = await ac.get(f"/api/v1/documents/{doc_id}/location/0")
        assert ctx_res.status_code == 200
        ctx_data = ctx_res.json()
        assert ctx_data["document_id"] == doc_id
        assert ctx_data["chunk"]["chunk_index"] == 0
        assert "anchor_label" in ctx_data["chunk"]
        assert "location_meta" in ctx_data["chunk"]

        # Clean up
        del_res = await ac.delete(f"/api/v1/documents/{doc_id}")
        assert del_res.status_code == 200
