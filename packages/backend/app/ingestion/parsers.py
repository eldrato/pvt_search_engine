import io
import os
import re
from abc import ABC, abstractmethod
from typing import List, Optional, Dict, Any
from pydantic import BaseModel, Field


class TextBlock(BaseModel):
    text: str
    page_number: Optional[int] = None
    line_start: Optional[int] = None
    line_end: Optional[int] = None
    paragraph_index: Optional[int] = None
    section_title: Optional[str] = None
    bbox: Optional[List[float]] = None  # [x0, y0, x1, y1] coordinates
    extra_meta: Dict[str, Any] = Field(default_factory=dict)


class ParsedDocument(BaseModel):
    title: str
    raw_text: str
    blocks: List[TextBlock] = Field(default_factory=list)
    mime_type: str
    metadata: Dict[str, Any] = Field(default_factory=dict)


class BaseParser(ABC):
    @abstractmethod
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        pass


class TextParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="replace")

        lines = text.splitlines()
        blocks = []
        para_idx = 0
        for i, line in enumerate(lines, 1):
            stripped = line.strip()
            if stripped:
                para_idx += 1
                blocks.append(TextBlock(
                    text=stripped,
                    line_start=i,
                    line_end=i,
                    paragraph_index=para_idx,
                    extra_meta={"paragraph": para_idx, "line": i}
                ))

        title = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()
        return ParsedDocument(
            title=title,
            raw_text=text,
            blocks=blocks,
            mime_type="text/plain",
            metadata={"filename": filename, "total_lines": len(lines), "total_paragraphs": para_idx}
        )


class PDFParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        import fitz  # PyMuPDF

        doc = fitz.open(stream=content, filetype="pdf")
        blocks: List[TextBlock] = []
        all_text_parts: List[str] = []

        title = doc.metadata.get("title") if doc.metadata else None
        if not title or not title.strip():
            title = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()

        total_paragraphs = 0
        for page_idx in range(len(doc)):
            page = doc[page_idx]
            page_num = page_idx + 1
            page_width = round(page.rect.width, 2)
            page_height = round(page.rect.height, 2)

            page_text = page.get_text("text")
            if page_text and page_text.strip():
                all_text_parts.append(page_text)

            # Extract structured blocks with bounding boxes: (x0, y0, x1, y1, "lines", block_no, block_type)
            raw_blocks = page.get_text("blocks")
            # Filter text blocks (block_type == 0)
            text_blocks = [b for b in raw_blocks if len(b) >= 7 and b[6] == 0 and b[4].strip()]

            if text_blocks:
                for idx, b in enumerate(text_blocks):
                    clean_text = b[4].strip()
                    bbox = [round(b[0], 2), round(b[1], 2), round(b[2], 2), round(b[3], 2)]
                    para_idx = idx + 1
                    total_paragraphs += 1
                    blocks.append(TextBlock(
                        text=clean_text,
                        page_number=page_num,
                        paragraph_index=para_idx,
                        bbox=bbox,
                        extra_meta={
                            "page": page_num,
                            "paragraph": para_idx,
                            "bbox": bbox,
                            "page_width": page_width,
                            "page_height": page_height
                        }
                    ))
            elif page_text and page_text.strip():
                # Fallback to paragraph splitting if raw blocks not available
                paras = [p.strip() for p in page_text.split("\n\n") if p.strip()]
                for idx, para in enumerate(paras):
                    para_idx = idx + 1
                    total_paragraphs += 1
                    blocks.append(TextBlock(
                        text=para,
                        page_number=page_num,
                        paragraph_index=para_idx,
                        extra_meta={"page": page_num, "paragraph": para_idx}
                    ))

        total_pages = len(doc)
        doc.close()

        raw_text = "\n\n".join(all_text_parts)
        return ParsedDocument(
            title=title.strip(),
            raw_text=raw_text,
            blocks=blocks,
            mime_type="application/pdf",
            metadata={
                "filename": filename,
                "total_pages": total_pages,
                "total_paragraphs": total_paragraphs
            }
        )


class DocxParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        from docx import Document as DocxDoc

        doc = DocxDoc(io.BytesIO(content))
        blocks: List[TextBlock] = []
        all_text: List[str] = []
        current_section = "Introduction"
        para_counter = 0

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Detect heading styles
            if para.style and para.style.name.startswith("Heading"):
                current_section = text
                continue

            para_counter += 1
            all_text.append(text)
            blocks.append(TextBlock(
                text=text,
                paragraph_index=para_counter,
                section_title=current_section,
                extra_meta={"section": current_section, "paragraph": para_counter}
            ))

        # Also parse table cells if present
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    para_counter += 1
                    all_text.append(row_text)
                    blocks.append(TextBlock(
                        text=row_text,
                        paragraph_index=para_counter,
                        section_title=current_section,
                        extra_meta={"is_table": True, "paragraph": para_counter}
                    ))

        raw_text = "\n\n".join(all_text)
        title = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()
        return ParsedDocument(
            title=title,
            raw_text=raw_text,
            blocks=blocks,
            mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            metadata={"filename": filename, "total_paragraphs": len(blocks)}
        )


class MarkdownParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="replace")

        lines = text.splitlines()
        blocks: List[TextBlock] = []
        current_heading = "General"
        first_h1 = None

        heading_pattern = re.compile(r"^(#{1,6})\s+(.+)$")

        current_block_lines = []
        start_line = 1

        for idx, line in enumerate(lines, 1):
            match = heading_pattern.match(line)
            if match:
                # Flush previous block
                if current_block_lines:
                    block_text = "\n".join(current_block_lines).strip()
                    if block_text:
                        blocks.append(TextBlock(
                            text=block_text,
                            line_start=start_line,
                            line_end=idx - 1,
                            section_title=current_heading,
                        ))
                    current_block_lines = []

                level = len(match.group(1))
                h_text = match.group(2).strip()
                if level == 1 and not first_h1:
                    first_h1 = h_text
                current_heading = h_text
                start_line = idx + 1
            else:
                if not current_block_lines:
                    start_line = idx
                current_block_lines.append(line)

        # Flush final block
        if current_block_lines:
            block_text = "\n".join(current_block_lines).strip()
            if block_text:
                blocks.append(TextBlock(
                    text=block_text,
                    line_start=start_line,
                    line_end=len(lines),
                    section_title=current_heading,
                ))

        title = first_h1 or os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()
        return ParsedDocument(
            title=title,
            raw_text=text,
            blocks=blocks,
            mime_type="text/markdown",
            metadata={"filename": filename, "total_lines": len(lines)}
        )


class CodeParser(BaseParser):
    EXTENSION_LANG_MAP = {
        ".py": "python",
        ".js": "javascript",
        ".jsx": "javascript",
        ".ts": "typescript",
        ".tsx": "typescript",
        ".go": "go",
        ".rs": "rust",
        ".java": "java",
        ".c": "c",
        ".cpp": "cpp",
        ".h": "c",
        ".hpp": "cpp",
        ".cs": "csharp",
        ".rb": "ruby",
        ".php": "php",
        ".css": "css",
        ".sql": "sql",
        ".json": "json",
        ".yaml": "yaml",
        ".yml": "yaml",
        ".sh": "bash",
        ".ps1": "powershell",
    }

    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        try:
            text = content.decode("utf-8")
        except UnicodeDecodeError:
            text = content.decode("latin-1", errors="replace")

        ext = os.path.splitext(filename)[1].lower()
        lang = self.EXTENSION_LANG_MAP.get(ext, "code")

        lines = text.splitlines()
        # Group code into logical windows of ~40 lines
        window_size = 40
        overlap = 10
        blocks: List[TextBlock] = []

        step = max(1, window_size - overlap)
        for i in range(0, len(lines), step):
            chunk_lines = lines[i:i + window_size]
            if not chunk_lines:
                continue
            block_text = "\n".join(chunk_lines)
            if block_text.strip():
                blocks.append(TextBlock(
                    text=block_text,
                    line_start=i + 1,
                    line_end=min(len(lines), i + len(chunk_lines)),
                    section_title=f"{filename}:{i + 1}-{min(len(lines), i + len(chunk_lines))}",
                    extra_meta={"language": lang, "filename": filename}
                ))

        return ParsedDocument(
            title=filename,
            raw_text=text,
            blocks=blocks,
            mime_type=f"text/x-{lang}",
            metadata={"filename": filename, "language": lang, "total_lines": len(lines)}
        )


class HTMLParser(BaseParser):
    def parse(self, content: bytes, filename: str) -> ParsedDocument:
        from bs4 import BeautifulSoup

        try:
            html = content.decode("utf-8")
        except UnicodeDecodeError:
            html = content.decode("latin-1", errors="replace")

        soup = BeautifulSoup(html, "html.parser")

        # Strip scripts, styles, and untrusted elements
        for tag in soup(["script", "style", "meta", "noscript", "iframe"]):
            tag.decompose()

        # Extract title
        title_tag = soup.find("title")
        title = title_tag.get_text().strip() if title_tag else ""
        if not title:
            h1 = soup.find("h1")
            title = h1.get_text().strip() if h1 else os.path.splitext(filename)[0]

        blocks: List[TextBlock] = []
        all_text = []

        # Find structural tags
        for el in soup.find_all(["h1", "h2", "h3", "h4", "p", "li", "pre", "blockquote"]):
            clean_text = el.get_text(separator=" ", strip=True)
            if clean_text:
                all_text.append(clean_text)
                blocks.append(TextBlock(
                    text=clean_text,
                    section_title=title,
                    extra_meta={"tag": el.name}
                ))

        raw_text = "\n\n".join(all_text)
        return ParsedDocument(
            title=title,
            raw_text=raw_text,
            blocks=blocks,
            mime_type="text/html",
            metadata={"filename": filename, "parsed_elements": len(blocks)}
        )


def get_parser(filename: str, mime_type: Optional[str] = None) -> BaseParser:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf" or mime_type == "application/pdf":
        return PDFParser()
    elif ext in [".docx"] or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return DocxParser()
    elif ext in [".md", ".markdown"]:
        return MarkdownParser()
    elif ext in [".html", ".htm"] or mime_type == "text/html":
        return HTMLParser()
    elif ext in CodeParser.EXTENSION_LANG_MAP or (mime_type and "code" in mime_type):
        return CodeParser()
    else:
        return TextParser()
